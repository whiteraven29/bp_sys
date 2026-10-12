"""The printed class attendance register (said 12 October 2026):

  * the college's template: the cover with the crest, the department, the
    academic year, the NTA level, the module code and name; signing pages of
    five sessions; the summary of sessions and the signatures; the footer;
  * exactly the sessions the tutor declared — 12 are pages of 5, 5 and 2, 45
    are nine pages — and the declaration is kept for next time;
  * the module's class as it stands, with NACTVET numbers: first sittings by
    name, then repeaters marked; a student taken off the module is not on it;
  * the module's tutors and the academic offices download it; another tutor
    does not.
"""

from datetime import date, timedelta
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.utils import timezone
from docx import Document
from rest_framework.test import APIClient

from . import registers
from .models import (
    AcademicYear, AdmissionWindow, ClassLevel, Department, Module, Programme, Semester, Student,
)
from .views import set_roles

User = get_user_model()


def office(username, *roles):
    user = User.objects.create_user(username, password='pw')
    set_roles(user, list(roles), full_name=username.title())
    client = APIClient()
    client.force_login(user)      # the download is a plain Django view: a real session
    return user, client


class RegisterTests(TestCase):
    def setUp(self):
        year = AcademicYear.objects.create(name='2026/2027', is_active=True)
        self.sem = Semester.objects.create(academic_year=year, number=1, is_active=True)
        level = ClassLevel.objects.create(name='NTA Level 5', order=5)
        pst = Programme.objects.create(
            department=Department.objects.create(name='Pharmaceutical Sciences', code='PST'),
            name='PST', code='PST')
        self.module = Module.objects.create(name='Pharmaceutical Microbiology', code='PST05101', teacher='T',
                                            class_level=level, semester=self.sem, programme=pst, credits=8)
        for reg, name, attempt in (('NS0003/0001/2024', 'Zawadi Mussa', Student.NORMAL),
                                   ('NS0001/0001/2023', 'Baraka Repeat', Student.REPEAT),
                                   ('NS0002/0001/2024', 'Amina Said', Student.NORMAL)):
            Student.objects.create(nactvet_reg_no=reg, name=name, module=self.module, attempt=attempt)
        self.tutor_user, self.tutor = office('tutor', 'tutor')
        self.module.teachers.add(self.tutor_user)

    def download(self, client=None, **params):
        return (client or self.tutor).get('/api/attendance/register/download/',
                                         {'module_id': self.module.id, **params})

    def document(self, **params):
        response = self.download(**params)
        self.assertEqual(response.status_code, 200, response.content[:200])
        return Document(BytesIO(response.content))

    @staticmethod
    def signing_tables(document):
        return [t for t in document.tables if t.rows[0].cells[0].text.startswith('Module Code:')]

    def test_the_cover(self):
        document = self.document(sessions=5)
        text = '\n'.join(p.text for p in document.paragraphs)
        for line in ('BLUE PHARMA INSTITUTE', 'OF SCIENCE AND TECHNOLOGY', 'DEPARTMENT OF PHARMACEUTICAL SCIENCES',
                     'ACADEMIC YEAR 2026/2027', 'CLASS ATTENDANCE REGISTER', 'FOR STUDENTS'):
            self.assertIn(line, text)
        cover = [[c.text for c in row.cells] for row in document.tables[0].rows]
        self.assertEqual(cover, [['NTA LEVEL', '5'], ['MODULE CODE', 'PST05101'],
                                 ['MODULE NAME', 'Pharmaceutical Microbiology']])
        self.assertEqual(len(document.inline_shapes), 1)                           # the crest
        section = document.sections[0]
        self.assertGreater(section.page_width, section.page_height)                  # landscape
        self.assertIn('Class Attendance Register 2026/2027', section.footer.paragraphs[0].text)

    def test_exactly_the_sessions_declared(self):
        tables = self.signing_tables(self.document(sessions=12))
        self.assertEqual(len(tables), 3)
        self.assertEqual([t.rows[0].cells[0].text.split('\t')[-1] for t in tables],
                         ['SESSIONS 1–5', 'SESSIONS 6–10', 'SESSIONS 11–12'])
        headings = [c.text for t in tables for c in t.rows[2].cells if c.text.startswith('SESSION')]
        self.assertEqual(headings, [f'SESSION {n}' for n in range(1, 13)])
        self.assertEqual(len(self.signing_tables(self.document(sessions=45))), 9)

    def test_the_summary_has_a_row_per_session(self):
        document = self.document(sessions=15)
        summary = next(t for t in document.tables if t.rows[0].cells[0].text.startswith('SESSION'))
        self.assertEqual([row.cells[0].text for row in summary.rows[1:]], [str(n) for n in range(1, 16)])
        signatures = document.tables[-1]
        self.assertEqual([c.text for c in signatures.rows[0].cells],
                         ['CLASS REPRESENTATIVE', 'MODULE FACILITATOR', 'HEAD OF DEPARTMENT (HoD)',
                          'HoD OFFICIAL STAMP'])

    def test_the_class_first_sittings_then_repeaters(self):
        table = self.signing_tables(self.document(sessions=5))[0]
        rows = [[c.text for c in row.cells[:3]] for row in table.rows[4:]]
        self.assertEqual(rows, [['1.', 'Amina Said', 'NS0002/0001/2024'],
                                ['2.', 'Zawadi Mussa', 'NS0003/0001/2024'],
                                ['3.', 'Baraka Repeat (Repeat)', 'NS0001/0001/2023']])

    def test_a_student_taken_off_the_module_is_not_on_it(self):
        Student.objects.filter(name='Zawadi Mussa').update(withdrawn_at=timezone.now())
        table = self.signing_tables(self.document(sessions=5))[0]
        self.assertEqual([row.cells[1].text for row in table.rows[4:]], ['Amina Said', 'Baraka Repeat (Repeat)'])

    def test_the_declaration_is_kept(self):
        response = self.tutor.post('/api/attendance/register/', {'module_id': self.module.id, 'sessions': 15},
                                   format='json')
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data['sessions'], 15)
        self.assertEqual(len(self.signing_tables(self.document())), 3)              # 15 without asking again
        self.assertIn('15_sessions', self.download()['Content-Disposition'])

    def test_without_a_declaration_the_templates_25(self):
        self.assertEqual(len(self.signing_tables(self.document())), 5)

    def test_a_sensible_number_only(self):
        for bad in (0, 101, 'many'):
            self.assertEqual(self.download(sessions=bad).status_code, 400)
            self.assertEqual(self.tutor.post('/api/attendance/register/', {
                'module_id': self.module.id, 'sessions': bad}, format='json').status_code, 400)

    def test_who_downloads(self):
        _, other = office('other', 'tutor')
        _, hod = office('hod', 'hod')
        self.assertEqual(self.download(other).status_code, 403)
        self.assertEqual(other.get('/api/attendance/register/', {'module_id': self.module.id}).status_code, 403)
        self.assertEqual(self.download(hod, sessions=5).status_code, 200)

    def test_it_says_while_admissions_are_open(self):
        AdmissionWindow.objects.create(semester=self.sem, opens_on=date.today() - timedelta(days=1),
                                       closes_on=date.today() + timedelta(days=7))
        info = self.tutor.get('/api/attendance/register/', {'module_id': self.module.id}).data
        self.assertTrue(info['admissions_open'])
        self.assertEqual((info['students'], info['repeaters']), (3, 1))

    def test_a_long_class_keeps_its_headings_on_every_page(self):
        table = self.signing_tables(self.document(sessions=5))[0]
        repeated = [row for row in table.rows[:4] if row._tr.trPr.find(
            '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblHeader') is not None]
        self.assertEqual(len(repeated), 4)

    def test_session_count_rules(self):
        self.assertEqual(registers.session_count(None), registers.DEFAULT_SESSIONS)
        self.module.register_sessions = 30
        self.assertEqual(registers.session_count('', self.module), 30)
        self.assertEqual(registers.session_count('8', self.module), 8)
