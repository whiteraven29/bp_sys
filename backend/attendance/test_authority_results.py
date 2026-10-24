"""The authority's grades, the results summary, and withholding results.

  * a D or an F from the authority is a fail — the module is repeated; with
    stars it is a supplementary;
  * an upload writes only to the semester it is for — a repeating student's
    earlier attempt is never touched;
  * only the examination officer (or the Principal) uploads; the Head of
    Department reads;
  * the summary gives each student's GPA and standing and each module's pass
    rate, from the same outcome the semester review uses;
  * withholding hides a student's results for a semester on the portal — they
    read "Withheld" — without changing them; the authority's *W* reads
    "Withheld" too.
"""

from django.contrib.auth import get_user_model
from django.test import TestCase
from rest_framework.test import APIClient

from . import progression
from .grading import parse_authority_grade
from .models import (
    AcademicYear, ClassLevel, Department, HeadOfDepartmentProfile, Module, Programme,
    ResultWithholding, SemesterReview, Semester, SemesterRegistration, Student, StudentProfile,
    StudentResult,
)
from .testkit import show_all_results

User = get_user_model()


class AuthorityBase(TestCase):
    def setUp(self):
        self.last_year = AcademicYear.objects.create(name='2025/2026')
        self.year = AcademicYear.objects.create(name='2026/2027', is_active=True)
        self.old_sem1 = Semester.objects.create(academic_year=self.last_year, number=1)
        self.sem1 = Semester.objects.create(academic_year=self.year, number=1, is_active=True)
        self.level4 = ClassLevel.objects.create(name='NTA Level 4', order=4)
        department = Department.objects.create(name='Pharmaceutical Sciences', code='PST')
        self.pst = Programme.objects.create(department=department, name='Pharmaceutical Sciences', code='PST')
        self.anatomy = self.module('PST04101', 'Human Anatomy')
        self.maths = self.module('PST04102', 'Pharmaceutical Calculations')

        self.exam_user = User.objects.create_superuser('exam', 'exam@x.com', 'pw')
        self.exam = APIClient(); self.exam.force_authenticate(self.exam_user)
        hod = User.objects.create_user('hod', password='pw', is_staff=True)
        HeadOfDepartmentProfile.objects.create(user=hod, full_name='HoD')
        self.hod = APIClient(); self.hod.force_authenticate(hod)

    def module(self, code, name, semester=None):
        return Module.objects.create(name=name, code=code, teacher='T', class_level=self.level4,
                                     semester=semester or self.sem1, programme=self.pst, credits=3)

    def student(self, reg_no, name, *modules):
        profile = StudentProfile.objects.create(nactvet_reg_no=reg_no, name=name)
        SemesterRegistration.objects.create(profile=profile, semester=self.sem1, programme=self.pst,
                                            class_level=self.level4, kind=SemesterRegistration.NEW)
        enrollments = [Student.objects.create(nactvet_reg_no=reg_no, name=name, profile=profile, module=module)
                       for module in modules or (self.anatomy, self.maths)]
        return profile, enrollments

    def upload(self, client, rows, semester=None):
        return client.post('/api/results/authority-grades/', {
            'rows': [{'reg_no': r, 'module_code': m, 'grade': g} for r, m, g in rows],
            'semester_id': (semester or self.sem1).id}, format='json')


class AuthorityGradeTests(AuthorityBase):
    def test_d_and_f_are_fails_and_stars_are_supplementaries(self):
        read = {grade: parse_authority_grade(grade, self.level4)['status'] for grade in
                ['A', 'B', 'C', 'D', 'F', 'C*', 'D*', 'F***', '*W*', 'FAIL']}

        self.assertEqual(read, {'A': 'PASS', 'B': 'PASS', 'C': 'PASS', 'D': 'FAIL', 'F': 'FAIL',
                                'C*': 'SUPP', 'D*': 'SUPP', 'F***': 'SUPP', '*W*': 'WITHHELD',
                                'FAIL': 'DISCONTINUED'})

    def test_an_f_from_the_authority_means_a_repeat(self):
        profile, _ = self.student('NIT/1', 'Asha')
        # A (4) and F (0): GPA 2.0, so a repeat rather than a discontinuation.
        self.upload(self.exam, [('NIT/1', 'PST04101', 'A'), ('NIT/1', 'PST04102', 'F')])

        outcome = progression.evaluate(profile, self.sem1)

        self.assertEqual(outcome['proposed'], SemesterReview.REPEAT)

    def test_the_upload_writes_only_to_the_semester_named(self):
        old_anatomy = self.module('PST04101', 'Human Anatomy', semester=self.old_sem1)
        profile, (repeat, _) = self.student('NIT/2', 'Baraka')
        first_attempt = Student.objects.create(nactvet_reg_no='NIT/2', name='Baraka', profile=profile,
                                               module=old_anatomy)
        StudentResult.objects.create(student=first_attempt, authority_grade='F', authority_status='FAIL',
                                     final_approved=True)

        response = self.upload(self.exam, [('NIT/2', 'PST04101', 'C')])

        self.assertEqual(response.data['saved'], 1, response.data)
        self.assertEqual(StudentResult.objects.get(student=repeat).authority_grade, 'C')
        self.assertEqual(StudentResult.objects.get(student=first_attempt).authority_grade, 'F')

    def test_the_semester_must_be_named(self):
        self.student('NIT/3', 'Chausiku')

        response = self.exam.post('/api/results/authority-grades/', {
            'rows': [{'reg_no': 'NIT/3', 'module_code': 'PST04101', 'grade': 'A'}]}, format='json')

        self.assertEqual(response.status_code, 400)

    def test_a_student_not_in_that_semester_is_reported_not_guessed(self):
        self.student('NIT/4', 'Daudi')

        response = self.upload(self.exam, [('NIT/4', 'PST04101', 'A')], semester=self.old_sem1)

        self.assertEqual(response.data['saved'], 0)
        self.assertIn('not enrolled', response.data['errors'][0])

    def test_the_head_of_department_reads_and_does_not_upload(self):
        self.student('NIT/5', 'Edina')

        refused = self.upload(self.hod, [('NIT/5', 'PST04101', 'A')])
        summary = self.hod.get(f'/api/results-summary/?semester_id={self.sem1.id}&class_level_id={self.level4.id}')

        self.assertEqual(refused.status_code, 403)
        self.assertEqual(summary.status_code, 200)


class ResultsSummaryTests(AuthorityBase):
    def setUp(self):
        super().setUp()
        self.student('NIT/10', 'Pass One')
        self.student('NIT/11', 'Supp Two')
        self.student('NIT/12', 'Repeat Three')
        self.student('NIT/13', 'Waiting Four')
        self.student('NIT/14', 'Disco Five')
        self.upload(self.exam, [
            ('NIT/10', 'PST04101', 'A'), ('NIT/10', 'PST04102', 'B'),
            ('NIT/11', 'PST04101', 'C'), ('NIT/11', 'PST04102', 'C*'),
            ('NIT/12', 'PST04101', 'A'), ('NIT/12', 'PST04102', 'F'),
            ('NIT/13', 'PST04101', 'B'),
            ('NIT/14', 'PST04101', 'B'), ('NIT/14', 'PST04102', 'F'),
        ])
        response = self.exam.get(f'/api/results-summary/?semester_id={self.sem1.id}&class_level_id={self.level4.id}')
        self.assertEqual(response.status_code, 200, response.data)
        self.data = response.data

    def test_how_many_passed_went_to_supp_or_repeat(self):
        counts = self.data['counts']

        self.assertEqual((counts['total'], counts['PASS'], counts['SUPP'], counts['REPEAT'], counts['DISCO'],
                          counts['INCOMPLETE']), (5, 1, 1, 1, 1, 1))

    def test_the_gpa_is_points_times_credits_over_credits(self):
        by_reg = {row['reg_no']: row for row in self.data['students']}

        # A (4) and B (3), three credits each: (12 + 9) / 6.
        self.assertEqual(by_reg['NIT/10']['gpa'], 3.5)
        self.assertEqual(by_reg['NIT/10']['classification'], 'First Class')
        # A (4) and F (0): 2.0 — the F is repeated.
        self.assertEqual((by_reg['NIT/12']['gpa'], by_reg['NIT/12']['remark']), (2.0, 'REPEAT'))
        # B (3) and F (0): 1.5 — below 2.0 is a discontinuation, not a repeat.
        self.assertEqual((by_reg['NIT/14']['gpa'], by_reg['NIT/14']['remark']), (1.5, 'DISCO'))
        self.assertEqual(by_reg['NIT/11']['supp'], ['PST04102'])

    def test_each_module_s_pass_rate(self):
        modules = {row['code']: row for row in self.data['modules']}

        # PST04101: five graded, all passed.
        self.assertEqual((modules['PST04101']['results'], modules['PST04101']['passed'],
                          modules['PST04101']['pass_rate']), (5, 5, 100.0))
        # PST04102: B, C*, F, F graded and one still waiting — one pass in four.
        self.assertEqual((modules['PST04102']['passed'], modules['PST04102']['supp'],
                          modules['PST04102']['failed'], modules['PST04102']['waiting'],
                          modules['PST04102']['pass_rate']), (1, 1, 2, 1, 25.0))


class WithholdingTests(AuthorityBase):
    def setUp(self):
        super().setUp()
        self.profile, (self.enrollment, _) = self.student('NIT/20', 'Fatuma')
        self.enrollment.set_portal_pin('Portal#2026', require_change=False)
        self.enrollment.save()
        self.upload(self.exam, [('NIT/20', 'PST04101', 'A'), ('NIT/20', 'PST04102', 'B')])
        session = self.client.session
        session['student_id'] = self.enrollment.id
        session.save()

    def statements(self):
        show_all_results()
        return self.client.get('/student-dashboard/').context['result_statements']

    def test_withheld_results_read_withheld_on_the_portal_and_are_not_changed(self):
        response = self.exam.post('/api/result-withholdings/', {
            'profile_id': self.profile.id, 'semester_id': self.sem1.id, 'reason': 'Fees not paid'}, format='json')

        self.assertEqual(response.status_code, 201, response.data)
        statement = self.statements()[0]
        self.assertTrue(statement['withheld'])
        self.assertEqual(statement['modules'], [])
        self.assertIsNone(statement['gpa'])
        self.assertContains(self.client.get('/student-dashboard/'), 'Results withheld')
        # The results stand, and the review still reads them.
        self.assertEqual(progression.evaluate(self.profile, self.sem1)['gpa'], 3.5)

    def test_releasing_shows_them_again_and_keeps_the_record(self):
        held = self.exam.post('/api/result-withholdings/', {
            'profile_id': self.profile.id, 'semester_id': self.sem1.id}, format='json').data

        self.exam.post(f'/api/result-withholdings/{held["id"]}/release/', {}, format='json')

        statement = self.statements()[0]
        self.assertFalse(statement['withheld'])
        self.assertEqual([row['grade'] for row in statement['modules']], ['A', 'B'])
        self.assertIsNotNone(ResultWithholding.objects.get(pk=held['id']).released_at)

    def test_only_the_examination_office_withholds(self):
        response = self.hod.post('/api/result-withholdings/', {
            'profile_id': self.profile.id, 'semester_id': self.sem1.id}, format='json')

        self.assertEqual(response.status_code, 403)

    def test_the_summary_shows_who_is_withheld(self):
        self.exam.post('/api/result-withholdings/', {
            'profile_id': self.profile.id, 'semester_id': self.sem1.id, 'reason': 'Fees'}, format='json')

        data = self.exam.get(f'/api/results-summary/?semester_id={self.sem1.id}&class_level_id={self.level4.id}').data

        self.assertEqual(data['counts']['withheld'], 1)
        self.assertEqual(data['students'][0]['withheld_reason'], 'Fees')

    def test_the_authority_s_star_w_reads_as_withheld(self):
        self.upload(self.exam, [('NIT/20', 'PST04102', '*W*')])

        grades = {row['code']: row['grade'] for row in self.statements()[0]['modules']}

        self.assertEqual(grades['PST04102'], 'Withheld')


class StarredGradesAtFaceValueTests(AuthorityBase):
    """NACTVET counts every grade at its own letter — a starred one too (B***
    is 3, D* is 1) — and cuts the GPA to one decimal. Checked on every 2024/25
    result sheet in diagrams/ (187 of 188 students)."""

    def gpa(self, profile):
        return progression.evaluate(profile, self.sem1)['gpa']

    def upload_supp(self, rows):
        return self.exam.post('/api/results/authority-grades/', {
            'rows': [{'reg_no': r, 'module_code': m, 'grade': g} for r, m, g in rows],
            'semester_id': self.sem1.id, 'sitting': 'supp'}, format='json')

    def test_a_starred_grade_counts_at_its_own_letter(self):
        level6 = ClassLevel.objects.create(name='NTA Level 6', order=6)

        points = {grade: parse_authority_grade(grade, self.level4)['points'] for grade in ['B***', 'C*', 'D*', 'F***']}

        self.assertEqual(points, {'B***': 3, 'C*': 2, 'D*': 1, 'F***': 0})
        self.assertEqual(parse_authority_grade('A*', level6)['points'], 5)

    def test_below_two_at_face_value_is_a_discontinuation_even_with_supplementaries(self):
        # Half failed (two of four) would go to supplementaries — but at face
        # value the GPA is 1.5, and NACTVET discontinues first.
        third, fourth = self.module('PST04103', 'Pharmacognosy'), self.module('PST04104', 'Pharmacology')
        profile, _ = self.student('NIT/30', 'Gift', self.anatomy, self.maths, third, fourth)
        self.upload(self.exam, [('NIT/30', 'PST04101', 'C'), ('NIT/30', 'PST04102', 'C'),
                                ('NIT/30', 'PST04103', 'D*'), ('NIT/30', 'PST04104', 'D*')])

        outcome = progression.evaluate(profile, self.sem1)

        self.assertEqual(outcome['gpa'], 1.5)
        self.assertEqual(outcome['proposed'], SemesterReview.DISCONTINUED)
        summary = self.exam.get(f'/api/results-summary/?semester_id={self.sem1.id}'
                                f'&class_level_id={self.level4.id}').data
        self.assertEqual(summary['students'][0]['remark'], 'DISCO')

    def test_a_nactvet_row_reads_as_nactvet_read_it(self):
        """Level 4, semester II 2024/25, NS2183/0046/2023: C C*** C F* B A over
        credits 4, 20, 12, 12, 12, 5 — NACTVET: GPA 1.9, FAIL. Counting the
        starred papers as C would have made it 2.3 and SUPP."""
        codes = [('PST04208', 4, 'C'), ('PST04209', 20, 'C***'), ('PST04210', 12, 'C'),
                 ('PST04211', 12, 'F*'), ('PST04212', 12, 'B'), ('PST04213', 5, 'A')]
        modules = [Module.objects.create(name=c, code=c, teacher='T', class_level=self.level4,
                                         semester=self.sem1, programme=self.pst, credits=credits)
                   for c, credits, _ in codes]
        profile, _ = self.student('NS2183/0046/2023', 'Sheet Row', *modules)
        self.upload(self.exam, [('NS2183/0046/2023', c, g) for c, _, g in codes])

        outcome = progression.evaluate(profile, self.sem1)

        self.assertEqual((outcome['gpa'], outcome['proposed']), (1.9, SemesterReview.DISCONTINUED))

    def test_the_gpa_is_cut_to_one_decimal(self):
        self.assertEqual([progression.cut_gpa(v) for v in (2.89, 3.47, 1.996, 2.0, 4.0)],
                         [2.8, 3.4, 1.9, 2.0, 4.0])

    def test_the_supplementary_sheet_is_its_own_sitting(self):
        passed, _ = self.student('NIT/31', 'Hope')
        failed, _ = self.student('NIT/32', 'Imani')
        self.upload(self.exam, [('NIT/31', 'PST04101', 'A'), ('NIT/31', 'PST04102', 'C*'),
                                ('NIT/32', 'PST04101', 'A'), ('NIT/32', 'PST04102', 'C*')])
        self.assertEqual((self.gpa(passed), self.gpa(failed)), (3.0, 3.0))       # (4 + 2) / 2

        # NACTVET's supplementary sheet re-lists every module: the passed paper
        # reads C, the one failed again is still starred.
        response = self.upload_supp([('NIT/31', 'PST04101', 'A'), ('NIT/31', 'PST04102', 'C'),
                                     ('NIT/32', 'PST04101', 'A'), ('NIT/32', 'PST04102', 'D***')])

        self.assertEqual((response.data['saved'], response.data['unchanged']), (2, 2))
        self.assertEqual(progression.evaluate(passed, self.sem1)['proposed'], SemesterReview.CLEAR)
        outcome = progression.evaluate(failed, self.sem1)
        self.assertEqual(outcome['proposed'], SemesterReview.REPEAT)
        self.assertEqual(outcome['gpa'], 2.5)                                 # (4 + 1) / 2
        # The first sitting's grade is never overwritten.
        result = StudentResult.objects.get(student__profile=failed, student__module=self.maths)
        self.assertEqual((result.authority_grade, result.authority_supp_grade), ('C*', 'D***'))

    def test_a_supplementary_sheet_cannot_change_a_module_that_was_not_one(self):
        profile, _ = self.student('NIT/34', 'Juma')
        self.upload(self.exam, [('NIT/34', 'PST04101', 'A'), ('NIT/34', 'PST04102', 'C*')])

        response = self.upload_supp([('NIT/34', 'PST04101', 'B'), ('NIT/34', 'PST04102', 'C')])

        self.assertEqual(response.data['saved'], 1)
        self.assertIn('was not a supplementary', response.data['errors'][0])
        result = StudentResult.objects.get(student__profile=profile, student__module=self.anatomy)
        self.assertEqual((result.authority_grade, result.authority_supp_grade), ('A', ''))

    def test_below_two_after_the_supplementary_is_a_discontinuation(self):
        # Level 4 supplementary sheet, NS3970/0014/2021: F* failed again, GPA 1.7, FAIL.
        profile, _ = self.student('NIT/35', 'Kassim')
        self.upload(self.exam, [('NIT/35', 'PST04101', 'B'), ('NIT/35', 'PST04102', 'C*')])
        self.assertEqual(progression.evaluate(profile, self.sem1)['proposed'], SemesterReview.PENDING)  # 2.5, supp
        self.upload_supp([('NIT/35', 'PST04102', 'F*')])

        outcome = progression.evaluate(profile, self.sem1)

        self.assertEqual((outcome['gpa'], outcome['proposed']), (1.5, SemesterReview.DISCONTINUED))

    def test_level_six_classes_are_nactvets(self):
        from .grading import gpa_classification
        level6 = ClassLevel.objects.create(name='NTA Level 6', order=6)
        # FIRST at 4.4 and UPPER SECOND at 4.2 are on the 2024/25 level 6 sheet;
        # LOWER SECOND at 2.8 on its supplementary sheet.
        self.assertEqual([gpa_classification(g, level6) for g in (4.4, 4.2, 3.5, 3.4, 2.8, 2.7, 2.6, 2.0)],
                         ['First Class', 'Upper Second', 'Upper Second', 'Lower Second',
                          'Lower Second', 'Lower Second', 'Pass', 'Pass'])

    def test_marks_entered_here_follow_the_same_rule(self):
        from .grading import grade_for_mark
        profile, (anatomy, maths) = self.student('NIT/33', 'Jabiri')
        StudentResult.objects.create(student=anatomy, assign1=16, assign2=16, cat1_theory=16, cat2_theory=16,
                                     end_theory=80, ca_approved=True, final_approved=True)
        # A failed end exam: supplementary, not yet sat — counted at its grade.
        result = StudentResult.objects.create(student=maths, assign1=16, assign2=16, cat1_theory=16,
                                              cat2_theory=16, end_theory=30, ca_approved=True,
                                              final_approved=True)

        rows = {row['code']: row for row in progression.evaluate(profile, self.sem1)['modules']}

        total = progression._result_serializer().get_final_total(result)
        self.assertEqual((rows['PST04102']['status'], rows['PST04102']['points']),
                         ('SUPP', grade_for_mark(total, self.level4)[1]))


class MoreThanHalfFailedTests(AuthorityBase):
    """The college's rule (3 October 2026): failing more than half of a
    semester's modules, the field module counted like any other, means no
    supplementary — the failed modules are repeated. Half or fewer go to
    supplementaries. A GPA below 2.0 (face value) is checked first and
    discontinues, as NACTVET does. The results are never changed."""

    def setUp(self):
        super().setUp()
        self.five = [self.anatomy, self.maths] + [
            self.module(f'PST0410{n}', name) for n, name in
            ((3, 'Pharmacognosy'), (4, 'Pharmacology'), (5, 'Field Practice'))]

    def grades(self, reg_no, *grades):
        self.upload(self.exam, [(reg_no, module.code, grade) for module, grade in zip(self.five, grades)])

    def test_three_of_five_failed_is_a_repeat_with_no_supplementary(self):
        profile, _ = self.student('NIT/50', 'Juma', *self.five)
        self.grades('NIT/50', 'A', 'A', 'C*', 'D*', 'F')

        outcome = progression.evaluate(profile, self.sem1)

        self.assertEqual(outcome['proposed'], SemesterReview.REPEAT)
        self.assertIn('more than half', outcome['reason'])
        self.assertEqual(outcome['gpa'], 2.2)                         # (4+4+2+1+0) / 5
        self.assertEqual(sorted(outcome['failed']), ['PST04103', 'PST04104', 'PST04105'])
        self.assertEqual(outcome['pending'], [])

    def test_below_two_is_a_discontinuation_before_more_than_half_is_counted(self):
        # NACTVET, level 5 semester II 2024/25: three or more failed with GPA 1.9
        # is FAIL; with GPA 2.2–2.4 it is REPEAT.
        profile, _ = self.student('NIT/51', 'Kemi', *self.five)
        self.grades('NIT/51', 'B', 'B', 'F', 'F', 'F')          # GPA 1.2

        outcome = progression.evaluate(profile, self.sem1)

        self.assertEqual(outcome['gpa'], 1.2)
        self.assertEqual(outcome['proposed'], SemesterReview.DISCONTINUED)

    def test_two_of_five_still_go_to_supplementaries(self):
        profile, _ = self.student('NIT/52', 'Lulu', *self.five)
        self.grades('NIT/52', 'B', 'B', 'B', 'C*', 'D*')

        self.assertEqual(progression.evaluate(profile, self.sem1)['proposed'], SemesterReview.PENDING)

    def test_three_of_four_is_a_repeat_and_two_of_four_is_not(self):
        three, _ = self.student('NIT/53', 'Mosi', *self.five[:4])
        two, _ = self.student('NIT/54', 'Neema', *self.five[:4])
        self.upload(self.exam, [('NIT/53', m.code, g) for m, g in zip(self.five, ['A', 'C*', 'C*', 'C*'])]
                    + [('NIT/54', m.code, g) for m, g in zip(self.five, ['A', 'A', 'C*', 'D*'])])

        self.assertEqual(progression.evaluate(three, self.sem1)['proposed'], SemesterReview.REPEAT)
        self.assertEqual(progression.evaluate(two, self.sem1)['proposed'], SemesterReview.PENDING)

    def test_confirming_raises_the_repeats_and_leaves_the_results_alone(self):
        from .models import OutstandingRepeat
        profile, enrollments = self.student('NIT/55', 'Omari', *self.five)
        self.grades('NIT/55', 'A', 'A', 'C*', 'D*', 'F')
        before = list(StudentResult.objects.filter(student__profile=profile)
                      .order_by('id').values_list('authority_grade', 'authority_status'))

        progression.build_reviews(self.sem1, actor=self.exam_user)
        review = SemesterReview.objects.get(profile=profile, semester=self.sem1)
        progression.confirm_review(review, SemesterReview.REPEAT, actor=self.exam_user)

        owed = set(OutstandingRepeat.objects.filter(profile=profile).values_list('module_code', flat=True))
        self.assertEqual(owed, {'PST04103', 'PST04104', 'PST04105'})
        after = list(StudentResult.objects.filter(student__profile=profile)
                     .order_by('id').values_list('authority_grade', 'authority_status'))
        self.assertEqual(before, after)

    def test_the_summary_and_the_declarations_screen_agree(self):
        self.student('NIT/56', 'Pili', *self.five)
        self.grades('NIT/56', 'A', 'A', 'C*', 'D*', 'F')

        summary = self.exam.get(f'/api/results-summary/?semester_id={self.sem1.id}'
                                f'&class_level_id={self.level4.id}').data
        row = summary['students'][0]
        self.assertEqual(row['remark'], 'REPEAT')
        self.assertEqual(row['supp'], [])
        module = self.exam.get(f'/api/results/module/?module_id={self.five[2].id}').data
        enrollment = Student.objects.get(nactvet_reg_no='NIT/56', module=self.five[2])
        self.assertEqual(module['direct_repeat'], [enrollment.id])

    def test_going_to_supplementaries_counts_them_at_their_own_letter(self):
        profile, _ = self.student('NIT/57', 'Rehema', *self.five)
        self.grades('NIT/57', 'B', 'B', 'B', 'D*', 'F*')

        outcome = progression.evaluate(profile, self.sem1)

        self.assertEqual(outcome['gpa'], 2.0)                         # (3+3+3+1+0) / 5
        self.assertEqual(outcome['proposed'], SemesterReview.PENDING)

    def test_failed_outright_counts_as_it_is(self):
        profile, _ = self.student('NIT/58', 'Saida', *self.five)
        self.grades('NIT/58', 'C', 'C', 'C', 'F', 'F')

        outcome = progression.evaluate(profile, self.sem1)

        self.assertEqual(outcome['gpa'], 1.2)                         # (2+2+2+0+0) / 5
        self.assertEqual(outcome['proposed'], SemesterReview.DISCONTINUED)


    def test_supplementaries_only_with_a_qualifying_gpa(self):
        """Below 2.0 at face value: discontinued now, not sent to supplementaries."""
        profile, _ = self.student('NIT/59', 'Tumaini', *self.five)
        self.grades('NIT/59', 'C', 'C', 'C', 'F', 'F*')

        outcome = progression.evaluate(profile, self.sem1)

        self.assertEqual(outcome['gpa'], 1.2)                         # (2+2+2+0+0) / 5
        self.assertEqual(outcome['proposed'], SemesterReview.DISCONTINUED)
        summary = self.exam.get(f'/api/results-summary/?semester_id={self.sem1.id}'
                                f'&class_level_id={self.level4.id}').data
        self.assertEqual(summary['students'][0]['remark'], 'DISCO')


class RepeatersSheetTests(AuthorityBase):
    """NACTVET keeps two sheets: continuing students, and repeaters — whose
    repeat grades go into the semester they failed, which then reads complete
    and cleared (said 12 October 2026)."""

    def setUp(self):
        super().setUp()
        # Last year: three modules; the student failed PST04103.
        self.old = {code: self.module(code, code, semester=self.old_sem1)
                    for code in ('PST04101', 'PST04102', 'PST04103')}
        self.profile = StudentProfile.objects.create(nactvet_reg_no='NIT/90', name='Rudia Repeater')
        SemesterRegistration.objects.create(profile=self.profile, semester=self.old_sem1, programme=self.pst,
                                            class_level=self.level4, kind=SemesterRegistration.NEW)
        self.old_sittings = {code: Student.objects.create(nactvet_reg_no='NIT/90', name='Rudia Repeater',
                                                          profile=self.profile, module=module)
                             for code, module in self.old.items()}
        self.upload(self.exam, [('NIT/90', 'PST04101', 'A'), ('NIT/90', 'PST04102', 'B'),
                                ('NIT/90', 'PST04103', 'D')], semester=self.old_sem1)
        # This year: the repeat sitting of PST04103 only, beside a continuing classmate.
        self.repeat_module = self.module('PST04103', 'Pharmacognosy')
        SemesterRegistration.objects.create(profile=self.profile, semester=self.sem1, programme=self.pst,
                                            class_level=self.level4, kind=SemesterRegistration.REPEATING)
        self.repeat = Student.objects.create(nactvet_reg_no='NIT/90', name='Rudia Repeater', profile=self.profile,
                                             module=self.repeat_module, attempt=Student.REPEAT,
                                             repeat_of=self.old_sittings['PST04103'])
        self.student('NIT/91', 'Cleo Continuing', self.anatomy, self.maths, self.repeat_module)
        self.upload(self.exam, [('NIT/91', 'PST04101', 'B'), ('NIT/91', 'PST04102', 'B'),
                                ('NIT/91', 'PST04103', 'B')])

    def summary(self, sheet):
        return self.exam.get(f'/api/results-summary/?semester_id={self.sem1.id}'
                             f'&class_level_id={self.level4.id}&sheet={sheet}').data

    def test_the_continuing_sheet_leaves_repeaters_out(self):
        self.upload(self.exam, [('NIT/90', 'PST04103', 'C')])
        self.assertEqual([s['reg_no'] for s in self.summary('continuing')['students']], ['NIT/91'])

    def test_a_passed_repeat_completes_the_failed_semester(self):
        self.upload(self.exam, [('NIT/90', 'PST04103', 'C')])
        row = self.summary('repeaters')['students'][0]
        self.assertEqual(row['origin'], self.old_sem1.label)
        self.assertEqual(row['repeats'], ['PST04103'])
        self.assertEqual(row['grades'], {'PST04101': 'A', 'PST04102': 'B', 'PST04103': 'C'})
        self.assertEqual((row['remark'], row['gpa']), ('PASS', 3.0))      # (4 + 3 + 2) / 3
        # The failed result is untouched.
        self.assertEqual(StudentResult.objects.get(student=self.old_sittings['PST04103']).authority_grade, 'D')

    def test_a_failed_repeat_is_a_discontinuation_on_the_sheet(self):
        self.upload(self.exam, [('NIT/90', 'PST04103', 'D*')])
        self.assertEqual(self.summary('repeaters')['students'][0]['remark'], 'DISCO')

    def test_waiting_until_the_repeat_grade_is_in(self):
        self.assertEqual(self.summary('repeaters')['students'][0]['remark'], 'INCOMPLETE')

    def test_the_portal_shows_the_completed_semester(self):
        self.upload(self.exam, [('NIT/90', 'PST04103', 'C')])
        self.repeat.set_portal_pin('Portal#2026', require_change=False)
        self.repeat.save()
        session = self.client.session
        session['student_id'] = self.repeat.id
        session.save()

        show_all_results()
        statements = self.client.get('/student-dashboard/').context['result_statements']

        self.assertEqual([s['label'] for s in statements], [self.old_sem1.label])
        old = statements[0]
        grades = {m['code']: (m['grade'], m.get('note', '')) for m in old['modules']}
        self.assertEqual(grades['PST04103'], ('C', f'Repeat, {self.sem1.label}'))
        self.assertEqual(grades['PST04101'][0], 'A')
        self.assertEqual(old['comment'], 'PASS')
        self.assertEqual(old['gpa'], 3.0)

    def test_a_withheld_repeater_is_marked_and_the_completed_semester_kept_back(self):
        """Said 24 October 2026: repeaters are withheld from their own sheet too,
        and withholding the repeat keeps back the semester it completes."""
        self.upload(self.exam, [('NIT/90', 'PST04103', 'C')])
        self.exam.post('/api/result-withholdings/', {'profile_ids': [self.profile.id], 'semester_id': self.sem1.id,
                                                     'reason': 'Fees'}, format='json')
        data = self.summary('repeaters')
        self.assertIsNotNone(data['students'][0]['withheld_id'])
        self.assertEqual(data['counts']['withheld'], 1)
        self.repeat.set_portal_pin('Portal#2026', require_change=False)
        self.repeat.save()
        session = self.client.session
        session['student_id'] = self.repeat.id
        session.save()
        show_all_results()
        statement = self.client.get('/student-dashboard/').context['result_statements'][0]
        self.assertEqual((statement['label'], statement['withheld'], statement['modules']),
                         (self.old_sem1.label, True, []))

    def test_a_repeater_sees_only_the_module_they_are_repeating(self):
        # Last time's first sittings are read for the statement, not listed as
        # this year's modules (said 13 October 2026).
        self.repeat.set_portal_pin('Portal#2026', require_change=False)
        self.repeat.save()
        session = self.client.session
        session['student_id'] = self.repeat.id
        session.save()
        show_all_results()
        page = self.client.get('/student-dashboard/')
        self.assertEqual([m['module_code'] for m in page.context['modules']], ['PST04103'])
        self.assertEqual(page.context['module_count'], 1)
        self.assertContains(page, 'Repeat</span>')
        self.assertEqual([s['label'] for s in page.context['result_statements']], [self.old_sem1.label])

    def test_the_ca_and_final_sheets_are_downloaded_apart(self):
        from io import BytesIO
        from openpyxl import load_workbook
        self.client.force_login(self.exam_user)
        # Opening the module's results page makes a result row for everybody
        # (as the page does before anything is downloaded).
        self.exam.get(f'/api/results/module/?module_id={self.repeat_module.id}')
        for url in ('/api/results/download/', '/api/results/download/final/'):
            for sitting, expected in (('continuing', {'NIT/91'}), ('repeaters', {'NIT/90'})):
                response = self.client.get(f'{url}?module_id={self.repeat_module.id}&sitting={sitting}')
                self.assertEqual(response.status_code, 200)
                sheet = load_workbook(BytesIO(response.content)).active
                found = {cell for row in sheet.iter_rows(values_only=True) for cell in row
                         if isinstance(cell, str) and cell.startswith('NIT/')}
                self.assertEqual(found, expected, f'{url} {sitting}')
                if sitting == 'repeaters':
                    self.assertIn('Repeaters', sheet.title)

    def test_each_result_row_says_its_sitting(self):
        rows = self.exam.get(f'/api/results/module/?module_id={self.repeat_module.id}').data['results']
        self.assertEqual({r['student_reg_no']: r['attempt'] for r in rows},
                         {'NIT/90': Student.REPEAT, 'NIT/91': Student.NORMAL})
