"""The printed class attendance register for one module (said 12 October 2026).

Built to the college's template (Class Attendance Register, 2026/2027): a cover
with the crest, the department, the academic year, the NTA level, the module
code and name; then the signing pages, five sessions to a page, 35 students to
a page with the heading rows repeated when a class runs longer; then the
summary of sessions and the signatures of the class representative, the
facilitator and the Head of Department.

The tutor declares how many sessions the module has (Module.register_sessions)
and the register has exactly that many — 15 sessions are three signing pages,
45 are nine. The students are the module's class as it stands: first sittings
by name, then repeaters, marked, as on the attendance screen.

Measurements are the template's own, read off it at 100 dpi: landscape Letter,
0.42" side margins, a 10.17" table, rows of 0.1875".
"""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .models import Student

COLLEGE = 'Blue Pharma Institute of Science and Technology'
COVER_NAME = ('BLUE PHARMA INSTITUTE', 'OF SCIENCE AND TECHNOLOGY')
LOGO = Path(__file__).resolve().parent / 'assets' / 'register_logo.png'

DEFAULT_SESSIONS = 25          # the template's own count
MAX_SESSIONS = 100
SESSIONS_PER_PAGE = 5
FONT = 'Arial'
SHADE = 'E6E6E6'

TABLE_WIDTH = 10.17
FIXED_COLUMNS = (0.42, 2.57, 1.66)                     # S/NO, name, registration number
SESSION_WIDTH = round((TABLE_WIDTH - sum(FIXED_COLUMNS)) / 5, 3)
SUMMARY_COLUMNS = (0.63, 1.25, 4.33, 1.32, 1.32, 1.32)
SIGNATURE_COLUMNS = (2.85, 2.87, 2.87, 1.58)
ROW = 0.1875
CELL_MARGIN = 72                                       # twips: 0.05"


class RegisterError(ValueError):
    """A register that cannot be made; the message says why."""


def session_count(value, module=None):
    """The number asked for, else the module's declared one, else the template's."""
    if value in (None, ''):
        return (module.register_sessions if module is not None and module.register_sessions
                else DEFAULT_SESSIONS)
    try:
        sessions = int(value)
    except (TypeError, ValueError):
        raise RegisterError('The number of sessions must be a whole number.')
    if not 1 <= sessions <= MAX_SESSIONS:
        raise RegisterError(f'Between 1 and {MAX_SESSIONS} sessions.')
    return sessions


def register_students(module):
    """The class as it stands: first sittings by name, then repeaters."""
    enrollments = list(Student.objects.studying().filter(module=module).order_by('name', 'nactvet_reg_no'))
    return ([s for s in enrollments if s.attempt != Student.REPEAT]
            + [s for s in enrollments if s.attempt == Student.REPEAT])


# ── low-level Word helpers ────────────────────────────────────────────────────

def _run(paragraph, text, *, size, bold=False, color=None):
    run = paragraph.add_run(text)
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return run


def _tight(paragraph, align=None):
    fmt = paragraph.paragraph_format
    fmt.space_before = fmt.space_after = Pt(0)
    fmt.line_spacing = 1.0
    if align is not None:
        paragraph.alignment = align
    return paragraph


def _shade(cell, fill=SHADE):
    props = cell._tc.get_or_add_tcPr()
    shading = OxmlElement('w:shd')
    shading.set(qn('w:val'), 'clear')
    shading.set(qn('w:color'), 'auto')
    shading.set(qn('w:fill'), fill)
    props.append(shading)


def _borders(cell, **sides):
    """Override a cell's borders: side=None removes it, side=size draws it."""
    props = cell._tc.get_or_add_tcPr()
    borders = props.find(qn('w:tcBorders'))
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        props.append(borders)
    for side, size in sides.items():
        edge = OxmlElement(f'w:{side}')
        if size is None:
            edge.set(qn('w:val'), 'nil')
        else:
            edge.set(qn('w:val'), 'single')
            edge.set(qn('w:sz'), str(size))
            edge.set(qn('w:color'), '000000')
        borders.append(edge)


def _row(row, height, *, header=False, exact=True):
    row.height = Inches(height)
    row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY if exact else WD_ROW_HEIGHT_RULE.AT_LEAST
    props = row._tr.get_or_add_trPr()
    props.append(OxmlElement('w:cantSplit'))
    if header:
        props.append(OxmlElement('w:tblHeader'))


def _table(document, widths, rows, align=WD_TABLE_ALIGNMENT.CENTER):
    """A fixed-layout grid table of exactly these column widths (inches)."""
    table = document.add_table(rows=rows, cols=len(widths))
    table.style = 'Table Grid'
    table.alignment = align
    table.autofit = False
    props = table._tbl.tblPr
    layout = OxmlElement('w:tblLayout')
    layout.set(qn('w:type'), 'fixed')
    props.append(layout)
    width = OxmlElement('w:tblW')
    width.set(qn('w:w'), str(int(sum(widths) * 1440)))
    width.set(qn('w:type'), 'dxa')
    for old in props.findall(qn('w:tblW')):
        props.remove(old)
    props.append(width)
    margins = OxmlElement('w:tblCellMar')
    for side, twips in (('top', 0), ('bottom', 0), ('left', CELL_MARGIN), ('right', CELL_MARGIN)):
        edge = OxmlElement(f'w:{side}')
        edge.set(qn('w:w'), str(twips))
        edge.set(qn('w:type'), 'dxa')
        margins.append(edge)
    props.append(margins)
    if align == WD_TABLE_ALIGNMENT.LEFT:
        # Word sets a left-aligned table out by its cell margin; line it up
        # with the centred tables' edge instead.
        indent = OxmlElement('w:tblInd')
        indent.set(qn('w:w'), str(CELL_MARGIN))
        indent.set(qn('w:type'), 'dxa')
        props.append(indent)
    for index, column in enumerate(table.columns):
        column.width = Inches(widths[index])
    for grid, w in zip(table._tbl.tblGrid.findall(qn('w:gridCol')), widths):
        grid.set(qn('w:w'), str(int(w * 1440)))
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index])
    return table


def _cell(cell, text='', *, size=7, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT,
          valign=WD_CELL_VERTICAL_ALIGNMENT.CENTER, color=None, shade=False):
    cell.vertical_alignment = valign
    paragraph = _tight(cell.paragraphs[0], align)
    for i, line in enumerate(str(text).split('\n')):
        if i:
            paragraph = _tight(cell.add_paragraph(), align)
        if line:
            _run(paragraph, line, size=size, bold=bold, color=color)
    if shade:
        _shade(cell)
    return cell


def _page_break(document):
    """A one-point paragraph that starts the next table on a new page."""
    paragraph = _tight(document.add_paragraph())
    paragraph.paragraph_format.page_break_before = True
    paragraph.paragraph_format.line_spacing = Pt(1)
    _run(paragraph, '', size=1)
    return paragraph


def _page_field(paragraph, size):
    for kind, text in (('begin', None), (None, 'PAGE'), ('separate', None), (None, '2'), ('end', None)):
        run = _run(paragraph, '', size=size)
        if kind:
            char = OxmlElement('w:fldChar')
            char.set(qn('w:fldCharType'), kind)
            run._r.append(char)
        elif text == 'PAGE':
            instr = OxmlElement('w:instrText')
            instr.set(qn('xml:space'), 'preserve')
            instr.text = ' PAGE '
            run._r.append(instr)
        else:
            run.text = text


# ── the register ──────────────────────────────────────────────────────────────

def _level_number(level):
    return str(level.order) if getattr(level, 'order', 0) else level.name


def _department(module):
    department = module.programme.department if module.programme_id and module.programme.department_id else None
    name = (department.name if department else 'Pharmaceutical Sciences').strip().upper()
    return name if name.startswith('DEPARTMENT') else f'DEPARTMENT OF {name}'


def _cover(document, module, year):
    if LOGO.exists():
        logo = _tight(document.add_paragraph(), WD_ALIGN_PARAGRAPH.CENTER)
        logo.paragraph_format.space_before = Pt(0)
        logo.add_run().add_picture(str(LOGO), width=Inches(1.9))
    name = _tight(document.add_paragraph(), WD_ALIGN_PARAGRAPH.CENTER)
    name.paragraph_format.space_before = Pt(14)
    _run(name, COVER_NAME[0], size=18, bold=True)
    name.add_run().add_break()
    _run(name, COVER_NAME[1], size=18, bold=True)

    rule = _tight(document.add_paragraph(), WD_ALIGN_PARAGRAPH.CENTER)
    rule.paragraph_format.space_before = Pt(12)
    rule.paragraph_format.left_indent = Inches(2.49)
    rule.paragraph_format.right_indent = Inches(2.48)
    rule.paragraph_format.line_spacing = Pt(2)
    borders = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    for key, value in (('val', 'single'), ('sz', '18'), ('space', '1'), ('color', '000000')):
        bottom.set(qn(f'w:{key}'), value)
    borders.append(bottom)
    rule._p.get_or_add_pPr().append(borders)

    department = _tight(document.add_paragraph(), WD_ALIGN_PARAGRAPH.CENTER)
    department.paragraph_format.space_before = Pt(14)
    _run(department, _department(module), size=10.5, bold=True)
    academic = _tight(document.add_paragraph(), WD_ALIGN_PARAGRAPH.CENTER)
    academic.paragraph_format.space_before = Pt(3)
    _run(academic, f'ACADEMIC YEAR {year}', size=10.5)

    title = _tight(document.add_paragraph(), WD_ALIGN_PARAGRAPH.CENTER)
    title.paragraph_format.space_before = Pt(22)
    _run(title, 'CLASS ATTENDANCE REGISTER', size=20, bold=True)
    sub = _tight(document.add_paragraph(), WD_ALIGN_PARAGRAPH.CENTER)
    _run(sub, 'FOR STUDENTS', size=10.5)
    _tight(document.add_paragraph()).paragraph_format.space_after = Pt(12)

    details = _table(document, (1.52, 5.14), 3)
    for row, (label, value) in zip(details.rows, (('NTA LEVEL', _level_number(module.class_level)),
                                                    ('MODULE CODE', module.code),
                                                    ('MODULE NAME', module.name))):
        _row(row, 0.435)
        _cell(row.cells[0], label, size=8.5, bold=True, shade=True)
        _cell(row.cells[1], value, size=8.5, bold=True)


def _signing_page(document, module, students, first, last):
    # Every session column is the template's width: a last page of two
    # sessions is the same page cut short, not two stretched columns.
    count = last - first + 1
    widths = FIXED_COLUMNS + (SESSION_WIDTH,) * count
    table = _table(document, widths, 4 + max(len(students), 1), align=WD_TABLE_ALIGNMENT.LEFT)

    # Row 0: the page's own heading, repeated with the others on a long class.
    heading = table.rows[0].cells[0].merge(table.rows[0].cells[-1])
    _row(table.rows[0], 0.20, header=True)
    _borders(heading, top=None, left=None, right=None)
    heading.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    line = _tight(heading.paragraphs[0])
    stops = line.paragraph_format.tab_stops
    stops.add_tab_stop(Inches(2.16))
    stops.add_tab_stop(Inches(sum(widths) - 0.1), WD_TAB_ALIGNMENT.RIGHT)
    _run(line, 'Module Code: ', size=7, bold=True)
    _run(line, module.code, size=7, bold=True)
    _run(line, '\tModule Name: ', size=7, bold=True)
    _run(line, module.name, size=7, bold=True)
    _run(line, f'\tSESSIONS {first}–{last}' if count > 1 else f'\tSESSION {first}', size=7, bold=True)

    # Rows 1–3: the column headings.
    _row(table.rows[1], 0.21, header=True, exact=count >= 4)    # fewer columns: the title wraps
    _row(table.rows[2], 0.22, header=True)
    _row(table.rows[3], 0.25, header=True)
    for index, label in enumerate(("S/NO", "STUDENT'S FULL NAME", 'NACTVET REG. NO')):
        merged = table.rows[1].cells[index].merge(table.rows[3].cells[index])
        _cell(merged, label, size=7, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, shade=True)
    title = table.rows[1].cells[3].merge(table.rows[1].cells[-1])
    # One column left: a smaller, tighter title, so even wrapped to four lines
    # the page still holds its 35 students.
    _cell(title, 'SIGNATURE OF STUDENTS FOR EACH SESSION IN A STATED DATE', size=7 if count > 1 else 5, bold=True,
          align=WD_ALIGN_PARAGRAPH.CENTER, shade=True)
    if count == 1:
        title.paragraphs[0].paragraph_format.line_spacing = Pt(5.6)
    for offset in range(count):
        _cell(table.rows[2].cells[3 + offset], f'SESSION {first + offset}', size=8.5, bold=True,
              align=WD_ALIGN_PARAGRAPH.CENTER)
        _cell(table.rows[3].cells[3 + offset], 'Date:', size=5.5, color='666666',
              valign=WD_CELL_VERTICAL_ALIGNMENT.BOTTOM)

    # One row per student.
    from .progression import provisional_profile_ids
    provisional = provisional_profile_ids()
    for number, row in enumerate(table.rows[4:], 1):
        _row(row, ROW)
        _cell(row.cells[0], f'{number}.', size=6.5)
        if number <= len(students):
            student = students[number - 1]
            name = student.name + (' (Repeat)' if student.attempt == Student.REPEAT else '') + (
                ' (Provisional)' if student.profile_id in provisional else '')
            _cell(row.cells[1], name, size=7)
            _cell(row.cells[2], student.nactvet_reg_no, size=7)
    return table


def _summary(document, sessions):
    title = _tight(document.add_paragraph(), WD_ALIGN_PARAGRAPH.CENTER)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(5)
    _run(title, 'SUMMARY OF SESSIONS', size=11, bold=True)

    table = _table(document, SUMMARY_COLUMNS, sessions + 1)
    _row(table.rows[0], 0.28, header=True)
    for cell, label in zip(table.rows[0].cells, ('SESSION\nNO.', 'DATE', 'SESSION DESCRIPTION / TOPIC COVERED',
                                                 'PRESENT', 'ABSENT', "FACILITATOR'S\nINITIALS")):
        _cell(cell, label, size=6.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, shade=True)
    for number, row in enumerate(table.rows[1:], 1):
        _row(row, 0.208)
        _cell(row.cells[0], str(number), size=6.5, align=WD_ALIGN_PARAGRAPH.CENTER)

    gap = _tight(document.add_paragraph())
    gap.paragraph_format.line_spacing = Pt(5)
    _run(gap, '', size=1)

    signatures = _table(document, SIGNATURE_COLUMNS, 2)
    _row(signatures.rows[0], 0.23)
    _row(signatures.rows[1], 1.04)
    for cell, label in zip(signatures.rows[0].cells, ('CLASS REPRESENTATIVE', 'MODULE FACILITATOR',
                                                      'HEAD OF DEPARTMENT (HoD)', 'HoD OFFICIAL STAMP')):
        _cell(cell, label, size=7, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, shade=True)
    for cell in signatures.rows[1].cells[:3]:
        _cell(cell, 'Name:\nSignature:\nDate:', size=7, bold=True, valign=WD_CELL_VERTICAL_ALIGNMENT.TOP)
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.line_spacing = Pt(12)


def build_register(module, sessions):
    """The register as a Word document."""
    year = module.semester.academic_year.name
    students = register_students(module)

    document = Document()
    normal = document.styles['Normal']
    normal.font.name = FONT
    normal.font.size = Pt(9)
    fonts = normal.element.get_or_add_rPr().get_or_add_rFonts()
    for key in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        fonts.set(qn(key), FONT)
    _tight(normal)

    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Inches(11), Inches(8.5)
    section.left_margin = section.right_margin = Inches(0.415)
    section.top_margin = Inches(0.33)
    section.bottom_margin = Inches(0.6)
    section.footer_distance = Inches(0.22)
    section.header_distance = Inches(0.2)
    section.different_first_page_header_footer = True     # the cover carries no footer

    footer = _tight(section.footer.paragraphs[0])
    stops = footer.paragraph_format.tab_stops
    for inherited in (3.25, 6.5):          # the Footer style's own centre and right stops
        stops.add_tab_stop(Inches(inherited), WD_TAB_ALIGNMENT.CLEAR)
    stops.add_tab_stop(Inches(TABLE_WIDTH), WD_TAB_ALIGNMENT.RIGHT)
    borders = OxmlElement('w:pBdr')
    top = OxmlElement('w:top')
    for key, value in (('val', 'single'), ('sz', '6'), ('space', '4'), ('color', '000000')):
        top.set(qn(f'w:{key}'), value)
    borders.append(top)
    footer._p.get_or_add_pPr().append(borders)
    _run(footer, f'{COLLEGE} | Class Attendance Register {year}', size=6, color='333333')
    _run(footer, '\tPage ', size=6, color='333333')
    _page_field(footer, 6)

    # The cover sits a little down the page, as on the template.
    spacer = _tight(document.add_paragraph())
    spacer.paragraph_format.line_spacing = Pt(22)
    _run(spacer, '', size=1)
    _cover(document, module, year)

    for first in range(1, sessions + 1, SESSIONS_PER_PAGE):
        _page_break(document)
        _signing_page(document, module, students, first, min(first + SESSIONS_PER_PAGE - 1, sessions))

    _page_break(document)
    _summary(document, sessions)
    end = _tight(document.add_paragraph())
    end.paragraph_format.line_spacing = Pt(1)
    _run(end, '', size=1)
    return document
